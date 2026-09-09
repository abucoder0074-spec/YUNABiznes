"""Read-only operational queries and Excel rendering from PostgreSQL data."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import case, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.enums import DocumentStatus, InventoryTransactionType
from app.models.catalogs import Product, Supplier, Unit, Warehouse
from app.models.inventory import InventoryTransaction, ProductionIssue
from app.models.purchasing import Payment, Purchase, PurchaseItem, SupplierLedger
from app.repositories.inventory import INBOUND_TRANSACTION_TYPES

ZERO = Decimal("0")


def _direction():
    return case((InventoryTransaction.transaction_type.in_(INBOUND_TRANSACTION_TYPES), 1), else_=-1)


def _decimal(value: object | None) -> Decimal:
    return Decimal(value or 0)


def _period(
    date_from: datetime | None, date_to: datetime | None, *, default_days: int = 30
) -> tuple[datetime, datetime]:
    end = date_to or datetime.now(UTC)
    start = date_from or end - timedelta(days=default_days)
    if start > end:
        raise ValueError("date_from must be earlier than or equal to date_to.")
    return start, end


async def inventory_balances(
    session: AsyncSession, *, warehouse_id: UUID | None = None, product_id: UUID | None = None
) -> list[dict[str, Any]]:
    quantity = func.coalesce(func.sum(InventoryTransaction.quantity * _direction()), ZERO).label("quantity")
    statement = (
        select(
            Warehouse.id.label("warehouse_id"),
            Warehouse.name.label("warehouse_name"),
            Product.id.label("product_id"),
            Product.name.label("product_name"),
            Unit.id.label("unit_id"),
            Unit.symbol.label("unit_symbol"),
            quantity,
        )
        .join(Warehouse, Warehouse.id == InventoryTransaction.warehouse_id)
        .join(Product, Product.id == InventoryTransaction.product_id)
        .join(Unit, Unit.id == InventoryTransaction.unit_id)
        .group_by(Warehouse.id, Warehouse.name, Product.id, Product.name, Unit.id, Unit.symbol)
        .having(quantity != ZERO)
        .order_by(Warehouse.name, Product.name, Product.id)
    )
    if warehouse_id is not None:
        statement = statement.where(InventoryTransaction.warehouse_id == warehouse_id)
    if product_id is not None:
        statement = statement.where(InventoryTransaction.product_id == product_id)
    rows = (await session.execute(statement)).mappings().all()
    return [{**row, "quantity": _decimal(row["quantity"])} for row in rows]


async def consumption(
    session: AsyncSession,
    *,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    warehouse_id: UUID | None = None,
    product_id: UUID | None = None,
) -> list[dict[str, Any]]:
    start, end = _period(date_from, date_to)
    consumed = func.coalesce(func.sum(InventoryTransaction.quantity), ZERO).label("consumed_quantity")
    statement = (
        select(
            Warehouse.id.label("warehouse_id"),
            Warehouse.name.label("warehouse_name"),
            Product.id.label("product_id"),
            Product.name.label("product_name"),
            Unit.id.label("unit_id"),
            Unit.symbol.label("unit_symbol"),
            consumed,
        )
        .join(Warehouse, Warehouse.id == InventoryTransaction.warehouse_id)
        .join(Product, Product.id == InventoryTransaction.product_id)
        .join(Unit, Unit.id == InventoryTransaction.unit_id)
        .where(
            InventoryTransaction.transaction_type == InventoryTransactionType.PRODUCTION_OUT,
            InventoryTransaction.occurred_at >= start,
            InventoryTransaction.occurred_at <= end,
        )
        .group_by(Warehouse.id, Warehouse.name, Product.id, Product.name, Unit.id, Unit.symbol)
        .order_by(Product.name, Warehouse.name, Product.id)
    )
    if warehouse_id is not None:
        statement = statement.where(InventoryTransaction.warehouse_id == warehouse_id)
    if product_id is not None:
        statement = statement.where(InventoryTransaction.product_id == product_id)
    days = max(Decimal("1"), Decimal(str((end - start).total_seconds() / 86400)))
    rows = (await session.execute(statement)).mappings().all()
    return [
        {
            **row,
            "consumed_quantity": _decimal(row["consumed_quantity"]),
            "average_daily_consumption": _decimal(row["consumed_quantity"]) / days,
        }
        for row in rows
    ]


async def replenishment_recommendations(
    session: AsyncSession,
    *,
    period_days: int = 30,
    warehouse_id: UUID | None = None,
    product_id: UUID | None = None,
) -> list[dict[str, Any]]:
    if not 1 <= period_days <= 365:
        raise ValueError("period_days must be between 1 and 365.")
    now = datetime.now(UTC)
    start = now - timedelta(days=period_days)
    stock_statement = select(
        InventoryTransaction.product_id,
        func.coalesce(func.sum(InventoryTransaction.quantity * _direction()), ZERO).label("quantity"),
    ).group_by(InventoryTransaction.product_id)
    use_statement = select(
        InventoryTransaction.product_id,
        func.coalesce(func.sum(InventoryTransaction.quantity), ZERO).label("quantity"),
    ).where(
        InventoryTransaction.transaction_type == InventoryTransactionType.PRODUCTION_OUT,
        InventoryTransaction.occurred_at >= start,
        InventoryTransaction.occurred_at <= now,
    ).group_by(InventoryTransaction.product_id)
    if warehouse_id is not None:
        stock_statement = stock_statement.where(InventoryTransaction.warehouse_id == warehouse_id)
        use_statement = use_statement.where(InventoryTransaction.warehouse_id == warehouse_id)
    stock = {row.product_id: _decimal(row.quantity) for row in (await session.execute(stock_statement))}
    usage = {row.product_id: _decimal(row.quantity) for row in (await session.execute(use_statement))}

    product_statement = select(Product, Unit).join(Unit, Unit.id == Product.unit_id).where(Product.is_active)
    if product_id is not None:
        product_statement = product_statement.where(Product.id == product_id)
    products = (await session.execute(product_statement.order_by(Product.name, Product.id))).all()
    result: list[dict[str, Any]] = []
    for product, unit in products:
        usable_stock = stock.get(product.id, ZERO)
        average_usage = usage.get(product.id, ZERO) / Decimal(period_days)
        lead_time = _decimal(product.lead_time_days)
        safety_stock = _decimal(product.safety_stock)
        reorder_point = average_usage * lead_time + safety_stock
        target_stock = _decimal(product.target_stock) if product.target_stock is not None else reorder_point
        recommended = max(target_stock - usable_stock, ZERO) if usable_stock <= reorder_point else ZERO
        reason = (
            f"Usable stock {usable_stock} {unit.symbol} is at or below the reorder point "
            f"{reorder_point} {unit.symbol} (average daily usage {average_usage}, "
            f"lead time {lead_time} days, safety stock {safety_stock} {unit.symbol})."
            if recommended > ZERO
            else f"Usable stock {usable_stock} {unit.symbol} is above the reorder point {reorder_point} {unit.symbol}."
        )
        result.append(
            {
                "product_id": product.id,
                "product_name": product.name,
                "unit_id": unit.id,
                "unit_symbol": unit.symbol,
                "usable_stock": usable_stock,
                "average_daily_usage": average_usage,
                "lead_time_days": lead_time,
                "safety_stock": safety_stock,
                "reorder_point": reorder_point,
                "target_stock": target_stock,
                "recommended_purchase_quantity": recommended,
                "reason": reason,
            }
        )
    return sorted(
        result,
        key=lambda item: (-item["recommended_purchase_quantity"], str(item["product_name"]).lower()),
    )


async def dashboard(session: AsyncSession) -> dict[str, Any]:
    now = datetime.now(UTC)
    today = now.replace(hour=0, minute=0, second=0, microsecond=0)
    balances = await inventory_balances(session)
    recommendations = await replenishment_recommendations(session)
    today_purchases = await session.scalar(
        select(func.coalesce(func.sum(PurchaseItem.total), ZERO))
        .join(Purchase, Purchase.id == PurchaseItem.purchase_id)
        .where(Purchase.status == DocumentStatus.POSTED, Purchase.purchased_at >= today)
    )
    today_consumption = await session.scalar(
        select(func.coalesce(func.sum(InventoryTransaction.quantity), ZERO)).where(
            InventoryTransaction.transaction_type == InventoryTransactionType.PRODUCTION_OUT,
            InventoryTransaction.occurred_at >= today,
        )
    )
    today_payments = await session.scalar(
        select(func.coalesce(func.sum(Payment.amount), ZERO)).where(
            Payment.status == DocumentStatus.POSTED, Payment.paid_at >= today
        )
    )
    supplier_debt = await session.scalar(
        select(func.coalesce(func.sum(SupplierLedger.debit_amount - SupplierLedger.credit_amount), ZERO))
    )
    recent_purchases = (
        await session.execute(
            select(Purchase.number, Supplier.name, Purchase.status, Purchase.purchased_at)
            .join(Supplier, Supplier.id == Purchase.supplier_id)
            .order_by(Purchase.created_at.desc(), Purchase.id)
            .limit(5)
        )
    ).all()
    recent_issues = (
        await session.execute(
            select(ProductionIssue.number, Warehouse.name, ProductionIssue.status, ProductionIssue.issued_at)
            .join(Warehouse, Warehouse.id == ProductionIssue.warehouse_id)
            .order_by(ProductionIssue.created_at.desc(), ProductionIssue.id)
            .limit(5)
        )
    ).all()
    recent_payments = (
        await session.execute(
            select(Payment.number, Supplier.name, Payment.status, Payment.amount, Payment.paid_at)
            .join(Supplier, Supplier.id == Payment.supplier_id)
            .order_by(Payment.created_at.desc(), Payment.id)
            .limit(5)
        )
    ).all()
    return {
        "current_stock_lines": len(balances),
        "current_stock_quantity": sum((row["quantity"] for row in balances), ZERO),
        "today_purchases": _decimal(today_purchases),
        "today_consumption": _decimal(today_consumption),
        "today_payments": _decimal(today_payments),
        "supplier_debt": _decimal(supplier_debt),
        "low_stock": [row for row in recommendations if row["recommended_purchase_quantity"] > ZERO][:10],
        "purchase_recommendations": [
            row for row in recommendations if row["recommended_purchase_quantity"] > ZERO
        ][:10],
        "recent_purchases": [
            {"number": number, "supplier_name": name, "status": status, "date": date}
            for number, name, status, date in recent_purchases
        ],
        "recent_production_issues": [
            {"number": number, "warehouse_name": name, "status": status, "date": date}
            for number, name, status, date in recent_issues
        ],
        "recent_payments": [
            {"number": number, "supplier_name": name, "status": status, "amount": amount, "date": date}
            for number, name, status, amount, date in recent_payments
        ],
    }


def _worksheet(workbook: Workbook, title: str, headers: list[str]):
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _finish(sheet) -> None:
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 50)
        sheet.column_dimensions[column[0].column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"


async def excel_report(
    session: AsyncSession,
    *,
    report_name: str,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
) -> tuple[str, bytes]:
    """Render a report from the same database queries used by the API, never from Excel state."""

    workbook = Workbook()
    start, end = _period(date_from, date_to)
    if report_name == "purchases":
        sheet = _worksheet(workbook, "Xaridlar", ["Sana", "Raqam", "Ta'minotchi", "Ombor", "Mahsulot", "Miqdor", "Narx", "Jami"])
        rows = await session.execute(
            select(Purchase.purchased_at, Purchase.number, Supplier.name, Warehouse.name, Product.name,
                   PurchaseItem.quantity, PurchaseItem.unit_price, PurchaseItem.total)
            .join(PurchaseItem, PurchaseItem.purchase_id == Purchase.id)
            .join(Supplier, Supplier.id == Purchase.supplier_id)
            .join(Warehouse, Warehouse.id == Purchase.warehouse_id)
            .join(Product, Product.id == PurchaseItem.product_id)
            .where(Purchase.status == DocumentStatus.POSTED, Purchase.purchased_at >= start, Purchase.purchased_at <= end)
            .order_by(Purchase.purchased_at, Purchase.number)
        )
        for row in rows:
            sheet.append(list(row))
        filename = "Xaridlar.xlsx"
    elif report_name == "payments":
        sheet = _worksheet(workbook, "Tolovlar", ["Sana", "Raqam", "Ta'minotchi", "Usul", "Summa", "Holat"])
        rows = await session.execute(
            select(Payment.paid_at, Payment.number, Supplier.name, Payment.method, Payment.amount, Payment.status)
            .join(Supplier, Supplier.id == Payment.supplier_id)
            .where(Payment.status == DocumentStatus.POSTED, Payment.paid_at >= start, Payment.paid_at <= end)
            .order_by(Payment.paid_at, Payment.number)
        )
        for row in rows:
            sheet.append(list(row))
        filename = "Tolovlar.xlsx"
    elif report_name == "debts":
        sheet = _worksheet(workbook, "Qarzlar", ["Ta'minotchi", "Debet", "Kredit", "Qarz"])
        rows = await session.execute(
            select(
                Supplier.name,
                func.coalesce(func.sum(SupplierLedger.debit_amount), ZERO),
                func.coalesce(func.sum(SupplierLedger.credit_amount), ZERO),
                func.coalesce(func.sum(SupplierLedger.debit_amount - SupplierLedger.credit_amount), ZERO),
            )
            .join(Supplier, Supplier.id == SupplierLedger.supplier_id)
            .group_by(Supplier.id, Supplier.name)
            .order_by(Supplier.name)
        )
        for row in rows:
            sheet.append(list(row))
        filename = "Qarzlar.xlsx"
    elif report_name == "inventory":
        sheet = _worksheet(workbook, "Ombor", ["Ombor", "Mahsulot", "Birlik", "Qoldiq"])
        for row in await inventory_balances(session):
            sheet.append([row["warehouse_name"], row["product_name"], row["unit_symbol"], row["quantity"]])
        filename = "Ombor.xlsx"
    elif report_name == "consumption":
        sheet = _worksheet(workbook, "Sarf", ["Ombor", "Mahsulot", "Birlik", "Sarf", "Kunlik o'rtacha"])
        for row in await consumption(session, date_from=start, date_to=end):
            sheet.append([row["warehouse_name"], row["product_name"], row["unit_symbol"], row["consumed_quantity"], row["average_daily_consumption"]])
        filename = "Sarf.xlsx"
    elif report_name == "recommendations":
        sheet = _worksheet(workbook, "Xarid tavsiyasi", ["Mahsulot", "Birlik", "Qoldiq", "Kunlik sarf", "Lead time", "Safety stock", "Reorder point", "Tavsiya", "Sabab"])
        for row in await replenishment_recommendations(session):
            sheet.append([row["product_name"], row["unit_symbol"], row["usable_stock"], row["average_daily_usage"], row["lead_time_days"], row["safety_stock"], row["reorder_point"], row["recommended_purchase_quantity"], row["reason"]])
        filename = "Xarid_tavsiyasi.xlsx"
    else:
        raise ValueError("Unknown report. Use purchases, payments, debts, inventory, consumption, or recommendations.")
    _finish(sheet)
    output = BytesIO()
    workbook.save(output)
    return filename, output.getvalue()
